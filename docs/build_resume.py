"""Generate a one-page Chinese resume.docx distilled from _pages/about.md.

Layout:
  - A4, narrow margins
  - Top header is a 2-column table:
      left  = name + affiliation + contact
      right = 2.5cm x 3.5cm bordered cell that reserves space for a 1-inch (一寸) photo
  - Sections: 个人简介 / 教育背景 / 学术论文 / 实习与学生工作 / 竞赛获奖 / 奖学金
  - Section titles use a colored left bar (▎) + bottom rule for a polished look.

Run: python docs/build_resume.py
Output: docs/resume.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PRIMARY = "00369F"
DARK = "1A1A1A"
GREY = "555555"
MUTED = "9A9A9A"
EAST_ASIA = "Microsoft YaHei"
LATIN = "Calibri"

OUTPUT = Path(__file__).resolve().parent / "resume.docx"


# ---------- low-level helpers ----------

def set_run_font(run, *, size=10.0, bold=False, italic=False, color=None,
                 east_asia=EAST_ASIA, latin=LATIN):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east_asia)


def add_bottom_border(paragraph, color=PRIMARY, size="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_cell_borders(cell, *, color="B0B0B0", size="6", sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        if side in sides:
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), size)
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), color)
        else:
            b.set(qn("w:val"), "nil")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def clear_cell_borders(cell):
    set_cell_borders(cell, sides=())


def set_cell_margins(cell, top=0.05, left=0.1, bottom=0.05, right=0.1):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(int(Cm(val).twips)))
        m.set(qn("w:type"), "dxa")
        tcMar.append(m)
    tcPr.append(tcMar)


def set_table_fixed_layout(table, col_widths_cm):
    """Force Word to honour the column widths instead of auto-fitting to content."""
    tbl = table._element
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total_twips = sum(int(Cm(w).twips) for w in col_widths_cm)
    tblW.set(qn("w:w"), str(total_twips))
    tblW.set(qn("w:type"), "dxa")

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    tblGrid = OxmlElement("w:tblGrid")
    for w in col_widths_cm:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(int(Cm(w).twips)))
        tblGrid.append(gridCol)
    tbl.insert(list(tbl).index(tblPr) + 1, tblGrid)


# ---------- content helpers ----------

def add_section_header(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    bar = p.add_run("▎ ")
    set_run_font(bar, size=13, bold=True, color=PRIMARY)
    title = p.add_run(text)
    set_run_font(title, size=12, bold=True, color=DARK)
    add_bottom_border(p)
    return p


def add_paragraph_runs(doc, segments, *, size=10, align=None, space_after=2,
                       line_spacing=1.2, indent_left=0.0, hanging=0.0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent_left:
        p.paragraph_format.left_indent = Cm(indent_left)
    if hanging:
        p.paragraph_format.first_line_indent = Cm(-hanging)
    for seg in segments:
        if isinstance(seg, str):
            text, opts = seg, {}
        else:
            text, opts = seg[0], (seg[1] if len(seg) > 1 else {})
        run = p.add_run(text)
        set_run_font(run, size=opts.get("size", size), bold=opts.get("bold", False),
                     italic=opts.get("italic", False), color=opts.get("color"))
    return p


def add_bullet(doc, segments, *, size=10, space_after=2):
    """segments: list of (text, {opts}) tuples."""
    full = [(("•  ", {"bold": True, "color": PRIMARY}))] + list(segments)
    return add_paragraph_runs(
        doc, full, size=size, space_after=space_after,
        line_spacing=1.22, indent_left=0.55, hanging=0.55,
    )


# ---------- builder ----------

def build_header(doc):
    # 一寸 photo is 2.5 cm × 3.5 cm (portrait). Total usable width ≈ 17.4 cm
    # at A4 with 1.8 cm side margins, so the info column takes the rest.
    info_w, photo_w, photo_h = 14.9, 2.5, 3.5

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False
    set_table_fixed_layout(table, [info_w, photo_w])

    info_cell, photo_cell = table.rows[0].cells
    info_cell.width = Cm(info_w)
    photo_cell.width = Cm(photo_w)

    row = table.rows[0]
    row.height = Cm(photo_h)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    # ---- left cell: name / role / contact ----
    clear_cell_borders(info_cell)
    set_cell_margins(info_cell, top=0.0, left=0.0, bottom=0.0, right=0.2)
    info_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    info_cell.paragraphs[0].text = ""  # clear default paragraph

    name_p = info_cell.paragraphs[0]
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(2)
    name_p.paragraph_format.line_spacing = 1.1
    name_run = name_p.add_run("章壹程")
    set_run_font(name_run, size=22, bold=True, color=DARK)
    sep_run = name_p.add_run("    ")
    set_run_font(sep_run, size=12)
    en_run = name_p.add_run("Yicheng Zhang")
    set_run_font(en_run, size=14, color=GREY)

    role_p = info_cell.add_paragraph()
    role_p.paragraph_format.space_before = Pt(0)
    role_p.paragraph_format.space_after = Pt(2)
    role_p.paragraph_format.line_spacing = 1.15
    role_run = role_p.add_run("南开大学 计算机学院  ·  计算机科学与技术  ·  本科在读")
    set_run_font(role_run, size=10.5, color=GREY)

    contact_p = info_cell.add_paragraph()
    contact_p.paragraph_format.space_before = Pt(0)
    contact_p.paragraph_format.space_after = Pt(0)
    contact_p.paragraph_format.line_spacing = 1.15
    contact_run = contact_p.add_run(
        "📮 3298809085@qq.com    📍 天津，中国    🔗 u2003yuge.github.io"
    )
    set_run_font(contact_run, size=10, color=GREY)

    # ---- right cell: 1-inch photo placeholder ----
    set_cell_borders(photo_cell, color="A8A8A8", size="6")
    set_cell_margins(photo_cell, top=0.0, left=0.0, bottom=0.0, right=0.0)
    photo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    photo_cell.paragraphs[0].text = ""
    placeholder = photo_cell.paragraphs[0]
    placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    placeholder.paragraph_format.space_before = Pt(0)
    placeholder.paragraph_format.space_after = Pt(0)
    line1 = placeholder.add_run("一寸照片")
    set_run_font(line1, size=10, color=MUTED)
    placeholder.add_run("\n")
    line2 = placeholder.add_run("2.5 × 3.5 cm")
    set_run_font(line2, size=8, color=MUTED, italic=True)


def build(target=OUTPUT):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.3)
        section.bottom_margin = Cm(1.3)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    build_header(doc)

    # spacer between header table and first section header
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    spacer_run = spacer.add_run("")
    set_run_font(spacer_run, size=4)

    # ---- 个人简介 ----
    add_section_header(doc, "个人简介")
    add_paragraph_runs(
        doc,
        [
            (
                "南开大学计算机学院计算机科学与技术专业本科在读。在校期间深耕专业领域，"
                "积极参与数学建模与算法类各类学科竞赛；以第二、第四作者身份发表两篇"
                "计算机视觉-遥感方向学术论文（AAAI 2026 Oral）。学生工作方面，曾任"
                "南开大学 ACM 算法协会社长 / 团支书，并担任数据库系统课程（计卓班）"
                "助教，积累了丰富的社团管理与教学辅助经验。",
                {"size": 10},
            ),
        ],
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.35,
        space_after=2,
    )

    # ---- 教育背景 ----
    add_section_header(doc, "教育背景")
    add_bullet(
        doc,
        [
            ("2023.06 – 至今    ", {"bold": True, "color": DARK}),
            ("南开大学 计算机学院，计算机科学与技术专业 本科在读（天津）。", {}),
        ],
    )

    # ---- 学术论文 ----
    add_section_header(doc, "学术论文")
    add_bullet(
        doc,
        [
            ("[AAAI 2026 Oral]  ", {"bold": True, "color": PRIMARY}),
            ("SM3Det: A Unified Model for Multi-Modal Remote Sensing Object Detection.", {}),
            ("   合作者：Yuxuan Li, Xiang Li, Yunheng Li, Yicheng Zhang (4th), Yimian Dai, Qibin Hou, Ming-Ming Cheng, Jian Yang.",
             {"size": 9, "color": GREY}),
        ],
    )
    add_bullet(
        doc,
        [
            ("[AAAI 2026 Oral]  ", {"bold": True, "color": PRIMARY}),
            ("Visual Instruction Pretraining for Domain-Specific Foundation Models.", {}),
            ("   合作者：Yuxuan Li, Yicheng Zhang (2nd), Wenhao Tang, Yimian Dai, Ming-Ming Cheng, Xiang Li, Jian Yang.",
             {"size": 9, "color": GREY}),
        ],
    )

    # ---- 实习与学生工作 ----
    add_section_header(doc, "实习与学生工作")
    add_bullet(
        doc,
        [
            ("2024 – 2026    ", {"bold": True, "color": DARK}),
            ("ReductLab 算法实习。", {}),
        ],
    )
    add_bullet(
        doc,
        [
            ("2026.02 – 2026.06    ", {"bold": True, "color": DARK}),
            ("南开大学 数据库系统课程（计卓班）助教，承担课程答疑、作业批改与实验指导。", {}),
        ],
    )
    add_bullet(
        doc,
        [
            ("2025.09 – 2026.06    ", {"bold": True, "color": DARK}),
            (
                "南开大学 ACM 算法协会 社长 / 团支书；任内社团获 2026 年度"
                "“五四红旗团支部标兵”、“先进社团”称号。",
                {},
            ),
        ],
    )
    add_bullet(
        doc,
        [
            ("2024 – 2026    ", {"bold": True, "color": DARK}),
            ("作为主要组织者，承办南开大学 ACM 新生赛与校赛 各 2 场。", {}),
        ],
    )
    add_bullet(
        doc,
        [
            ("2025.04    ", {"bold": True, "color": DARK}),
            ("联合举办首届 津冀联合高校大学生程序设计竞赛。", {}),
        ],
    )

    # ---- 竞赛获奖 ----
    add_section_header(doc, "竞赛获奖")
    add_bullet(
        doc,
        [
            ("ACM-ICPC 亚洲区域赛  ", {"bold": True, "color": DARK}),
            (
                "银牌 ×3（2025 上海 / 2024 杭州 / 2024 成都）；铜牌 ×3（2025 南京 / "
                "2023 西安 / 2023 合肥）；2024 ICPC 中国陕西省邀请赛 金牌。",
                {},
            ),
        ],
    )
    add_bullet(
        doc,
        [
            ("数学建模  ", {"bold": True, "color": DARK}),
            (
                "2025 全国大学生数学建模竞赛 天津赛区 一等奖；2024 同竞赛 二等奖；"
                "2025 美国大学生数学建模竞赛（MCM/ICM）Honorable Mention。",
                {},
            ),
        ],
    )
    add_bullet(
        doc,
        [
            ("程序设计 / 系统能力  ", {"bold": True, "color": DARK}),
            (
                "2025 第七届码蹄杯全国大学生程序设计竞赛国赛 金奖；2025 全国大学生"
                "计算机系统能力大赛（先导杯）三等奖。",
                {},
            ),
        ],
    )
    add_bullet(
        doc,
        [
            ("其他  ", {"bold": True, "color": DARK}),
            (
                "2025 南开大学“火山杯”AI 应用创新大赛 第一名；"
                "2024 CCF 计算机软件能力认证（CSP）成绩位列全国前 1%。",
                {},
            ),
        ],
    )

    # ---- 奖学金 ----
    add_section_header(doc, "奖学金")
    add_bullet(
        doc,
        [
            ("2025    ", {"bold": True, "color": DARK}),
            ("南开大学 创新奖学金。", {}),
        ],
    )
    add_bullet(
        doc,
        [
            ("2024    ", {"bold": True, "color": DARK}),
            ("南开大学 创新奖学金。", {}),
        ],
    )

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(f"wrote {target}")


if __name__ == "__main__":
    build()
